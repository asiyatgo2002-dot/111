"""Сохранение результатов в Excel."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import Product
from .scoring import ScoringConfig

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MP_FILL = {"Wildberries": PatternFill("solid", fgColor="F3E5F5"), "Ozon": PatternFill("solid", fgColor="E3F2FD")}
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLUMNS = [
    # (заголовок, ширина, формат, функция)
    ("№", 5, None, None),
    ("Маркетплейс", 13, None, lambda p: p.marketplace),
    ("Товар", 55, None, lambda p: p.name),
    ("Бренд", 16, None, lambda p: p.brand),
    ("Цена, ₽", 11, "#,##0", lambda p: p.price),
    ("Цена до скидки, ₽", 13, "#,##0", lambda p: p.old_price),
    ("Скидка, %", 9, "0.0", lambda p: p.discount_pct),
    ("Рейтинг", 8, "0.0", lambda p: p.rating),
    ("Отзывов", 10, "#,##0", lambda p: p.reviews),
    ("Продажи", 11, "#,##0", lambda p: p.sales),
    ("Источник продаж", 26, None, lambda p: p.sales_source),
    ("Балл продаж", 9, "0.0", lambda p: p.sales_score),
    ("Балл цена/качество", 11, "0.0", lambda p: p.value_score),
    ("Балл качества", 9, "0.0", lambda p: p.quality_score),
    ("Итоговый балл", 10, "0.0", lambda p: p.total_score),
    ("Артикул", 12, None, lambda p: p.sku),
    ("Ссылка", 14, None, lambda p: p.url),
    ("Комментарий ИИ", 45, None, lambda p: p.comment),
]


def _write_table(ws, products: list[Product], start_row: int = 1) -> None:
    for col, (title, width, _, _) in enumerate(COLUMNS, 1):
        c = ws.cell(row=start_row, column=col, value=title)
        c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[start_row].height = 32

    for i, p in enumerate(products, 1):
        row = start_row + i
        for col, (title, _, fmt, getter) in enumerate(COLUMNS, 1):
            value = i if getter is None else getter(p)
            c = ws.cell(row=row, column=col, value=value)
            c.border = BORDER
            c.fill = MP_FILL.get(p.marketplace, PatternFill())
            c.alignment = Alignment(vertical="top", wrap_text=title in ("Товар", "Комментарий ИИ"))
            if fmt:
                c.number_format = fmt
            if title == "Ссылка" and value:
                c.hyperlink = value
                c.value = "открыть"
                c.font = Font(color="0563C1", underline="single")

    last = start_row + len(products)
    if products:
        total_col = get_column_letter(next(i for i, c in enumerate(COLUMNS, 1) if c[0] == "Итоговый балл"))
        ws.conditional_formatting.add(
            f"{total_col}{start_row + 1}:{total_col}{last}",
            ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile",
                           mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"),
        )
    ws.freeze_panes = ws.cell(row=start_row + 1, column=4)
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(COLUMNS))}{last}"


def save_report(
    path: Path,
    topic: str,
    queries: list[str],
    top: list[Product],
    all_ranked: list[Product],
    cfg: ScoringConfig,
    summary: str = "",
    notes: list[str] | None = None,
    generated_at: datetime | None = None,
    top_n: int | None = None,
) -> Path:
    generated_at = generated_at or datetime.now()
    top_n = top_n or len(top)
    wb = Workbook()

    ws = wb.active
    ws.title = f"Топ-{top_n}"
    ws["A1"] = f"Топ-{top_n}: «{topic}» — Ozon и Wildberries"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Актуально на {generated_at:%d.%m.%Y %H:%M}. Сортировка по итоговому баллу (продажи + цена/качество)."
    ws["A2"].font = Font(italic=True, color="595959")
    _write_table(ws, top, start_row=4)

    for mp in ("Wildberries", "Ozon"):
        items = [p for p in all_ranked if p.marketplace == mp]
        if items:
            _write_table(wb.create_sheet(mp), items)

    _write_table(wb.create_sheet("Все кандидаты"), all_ranked)

    info = wb.create_sheet("Методика")
    info.column_dimensions["A"].width = 30
    info.column_dimensions["B"].width = 110
    rows = [
        ("Тема", topic),
        ("Поисковые запросы", "; ".join(queries)),
        ("Дата сбора", f"{generated_at:%d.%m.%Y %H:%M}"),
        ("Кандидатов после фильтров", len(all_ranked)),
        ("Фильтры", f"рейтинг ≥ {cfg.min_rating}, отзывов ≥ {cfg.min_reviews}, "
                    f"цена ≥ {cfg.min_price_ratio:.0%} от медианы (отсев аксессуаров)"),
        ("Итоговый балл", f"{cfg.w_sales:.0%} продажи + {cfg.w_value:.0%} цена/качество + "
                          f"{cfg.w_quality:.0%} качество (все баллы 0–100 внутри выборки)"),
        ("Продажи", "Wildberries — счётчик заказов карточки товара; Ozon — бейдж «купили» в выдаче. "
                    f"Если маркетплейс число не показывает — оценка: отзывы ÷ {cfg.review_rate:.0%} "
                    "(доля покупателей, оставляющих отзыв). Продажи по шкале логарифмические."),
        ("Качество", f"Байесовский рейтинг: оценка, сглаженная {cfg.prior_reviews} «виртуальными» отзывами "
                     "со средней оценкой выборки — мало отзывов = меньше доверия."),
        ("Цена/качество", "Байесовский рейтинг, делённый на 1 + ln(цена / медиана + 1)."),
    ]
    if summary:
        rows.append(("Вывод ИИ-аналитика", summary))
    for note in notes or []:
        rows.append(("Примечание", note))
    for r, (k, v) in enumerate(rows, 1):
        info.cell(row=r, column=1, value=k).font = Font(bold=True)
        cell = info.cell(row=r, column=2, value=v)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
